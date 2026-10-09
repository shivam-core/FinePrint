# app/extraction.py
import io
from pypdf import PdfReader
import docx

def extract_txt(file_bytes: bytes) -> tuple[str, list[str]]:
    warnings = []
    try:
        # Try UTF-8 first
        text = file_bytes.decode('utf-8-sig')
    except UnicodeDecodeError:
        try:
            # Fallback to latin-1 or similar if decode fails? 
            # Requirements say "if decode fails return an encoding instruction instead of corrupting characters."
            return "", ["Unable to read text file. Please ensure the file is saved with UTF-8 encoding."]
        except Exception as e:
            return "", [f"Error reading text: {str(e)}"]
    
    return text, warnings

def extract_pdf(file_bytes: bytes) -> tuple[str, int, list[str]]:
    warnings = []
    text_parts = []
    
    try:
        reader = PdfReader(io.BytesIO(file_bytes))
        if reader.is_encrypted:
            return "", 0, ["PDF is encrypted and requires a password. Cannot extract text."]
            
        num_pages = len(reader.pages)
        empty_pages = []
        for i, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if not page_text or not page_text.strip():
                empty_pages.append(i + 1)
            else:
                text_parts.append(page_text)
                
        if empty_pages:
            warnings.append(f"Pages {', '.join(map(str, empty_pages))} have no extractable text.")
            
        full_text = "\n\n".join(text_parts).strip()
        if not full_text:
            warnings.append("This PDF appears scanned or contains no extractable text. Upload a text-based PDF or paste the text. OCR is not included.")
            
        return full_text, num_pages, warnings
    except Exception as e:
        return "", 0, [f"Error reading PDF: {str(e)}"]

def extract_docx(file_bytes: bytes) -> tuple[str, list[str]]:
    warnings = []
    text_parts = []
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        # Walk document body in order (paragraphs and tables)
        for element in doc.element.body:
            if element.tag.endswith('p'):
                # it's a paragraph
                for p in doc.paragraphs:
                    if p._element == element:
                        text_parts.append(p.text)
                        break
            elif element.tag.endswith('tbl'):
                # it's a table
                for t in doc.tables:
                    if t._element == element:
                        for row in t.rows:
                            row_text = []
                            for cell in row.cells:
                                row_text.append(cell.text.strip())
                            text_parts.append(" | ".join(row_text))
                        break
                        
        warnings.append("Header/footer/comment content is excluded from extraction.")
        full_text = "\n\n".join(text_parts).strip()
        return full_text, warnings
    except Exception as e:
        return "", [f"Error reading DOCX: {str(e)}"]
